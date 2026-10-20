"""账号台账（xlsx）：一行一个邀请码账号（2026-10-20 立）

手工服务期用它看「三个账号各自用成什么样」。这组用例钉三条，任何一条回退都会让台账说谎：

1. **作废的事件不算** —— 曾出现的错：自己逐行 json.loads 读账本，把 void 掉的账
   也算进合计，支出立刻虚高；
2. **同号发票只算一次** —— 与工作台同一口径；
3. **重新生成不抹掉手填列** —— 反馈/是否付费是几周攒下来的，刷新一次就没了等于白记。
"""
import io
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import make_tenant_ledger as ledger  # noqa: E402


def _mk(root: Path, slug: str, acts: list[dict]) -> None:
    d = root / "tenants" / slug
    d.mkdir(parents=True, exist_ok=True)
    (d / "finance_events.jsonl").write_text(
        "\n".join(json.dumps(a, ensure_ascii=False) for a in acts) + "\n", encoding="utf-8")


def _env(tmp_path: Path) -> Path:
    rows = [
        {"slug": "aaaaaaaa", "label": "账号一", "active": True,
         "created_at": "2026-10-19T10:00:00", "last_seen": "2026-10-20T09:00:00"},
        {"slug": "bbbbbbbb", "label": "账号二", "active": False,
         "created_at": "2026-10-19T11:00:00", "last_seen": ""},
    ]
    (tmp_path / "tenants.json").write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")
    _mk(tmp_path, "aaaaaaaa", [
        {"event_id": "e1", "type": "income", "amount_cents": 1000000, "ts": "2026-10-20T10:00:00",
         "category": "接单", "dedupe_key": "k1"},
        {"event_id": "e2", "type": "expense", "amount_cents": 96200, "ts": "2026-10-20T10:05:00",
         "category": "住宿", "dedupe_key": "k2", "extra": {"invoice_no": "INV1"}},
        {"event_id": "v1", "type": "void", "amount_cents": 0, "ts": "2026-10-20T10:10:00",
         "targets": ["e2"], "dedupe_key": "k-void"},
        {"event_id": "e3", "type": "expense", "amount_cents": 96200, "ts": "2026-10-20T10:11:00",
         "category": "住宿", "dedupe_key": "k3", "extra": {"invoice_no": "INV1"}},
        {"event_id": "e4", "type": "expense", "amount_cents": 96200, "ts": "2026-10-20T10:12:00",
         "category": "住宿", "dedupe_key": "k4", "extra": {"invoice_no": "INV1"}},
    ])
    return tmp_path


def _row(path: Path, name: str) -> dict:
    """按表头取一行的字典 —— 列序会变（后来加入了「邀请码」列），用下标会脆。"""
    wb = load_workbook(path)
    ws = wb["账号台账"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == name:
            return dict(zip(head, row))
    raise AssertionError(f"台账里没有 {name}")


def test_voided_event_is_not_counted(tmp_path):
    root = _env(tmp_path)
    out = ledger.build(root, tmp_path / "ledger.xlsx")
    row = _row(out, "账号一")
    # 收入 10,000；支出只有一笔 962（被 void 的那笔不算，重复发票也只算一次）
    assert row["收入(元)"] == 10000
    assert row["支出(元)"] == 962
    assert row["净额(元)"] == 9038
    assert row["账本笔数"] == 2


def test_duplicate_invoice_counted_once(tmp_path):
    root = _env(tmp_path)
    out = ledger.build(root, tmp_path / "ledger.xlsx")
    row = _row(out, "账号一")
    assert row["支出(元)"] == 962        # 两笔同号票只算一次


def test_inactive_account_is_marked_and_empty(tmp_path):
    root = _env(tmp_path)
    out = ledger.build(root, tmp_path / "ledger.xlsx")
    row = _row(out, "账号二")
    assert row["状态"] == "停用"
    assert row["收入(元)"] == 0 and row["支出(元)"] == 0


def test_regenerate_keeps_manual_columns(tmp_path):
    """重新生成时按账号名保留手填的反馈/是否付费/备注。"""
    root = _env(tmp_path)
    out = tmp_path / "ledger.xlsx"
    ledger.build(root, out)

    wb = load_workbook(out)
    ws = wb["账号台账"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "账号一":
            row[head.index("反馈 / 用得怎么样")].value = "用了一周，说对账比原来快"
            row[head.index("是否付费")].value = "已付 99"
            row[head.index("备注")].value = "下月回访"
    wb.save(out)

    ledger.build(root, out)          # 再生成一次
    row = _row(out, "账号一")
    assert row["反馈 / 用得怎么样"] == "用了一周，说对账比原来快"
    assert row["是否付费"] == "已付 99"
    assert row["备注"] == "下月回访"
    assert row["收入(元)"] == 10000   # 自动列仍然是最新的


def test_invite_code_column_comes_from_the_codes_file(tmp_path):
    """邀请码列：服务器上没有明文码，只能由 --codes 指定的本地文件填。"""
    root = _env(tmp_path)
    codes_file = tmp_path / "codes.json"
    codes_file.write_text(json.dumps({"账号一": "CL-AAAA-BBBB-CCCC"}, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "ledger.xlsx"
    ledger.build(root, out, codes={"账号一": "CL-AAAA-BBBB-CCCC"})
    assert _row(out, "账号一")["邀请码"] == "CL-AAAA-BBBB-CCCC"
    assert _row(out, "账号二")["邀请码"] in (None, "")   # 没给的留空，不编（空串读回来是 None）
