#!/usr/bin/env python3
"""生成「账号台账」Excel：一行一个邀请码账号（2026-10-20 立）

## 这是什么

手工服务期的管理台账。三个账号（你 + 两个朋友）一人一行，把**能自动算的**从账本里
算出来（笔数 / 收支 / 净额 / 最近使用），把**只能人来判断的**留空给你手填
（反馈、是否付费、备注）。

## 三条纪律

1. **只读**。只读 `tenants.json` 与各租户的 `finance_events.jsonl`，不写、不建、不改
   任何数据文件 —— 台账是观察工具，不是操作工具。
2. **口径与工作台一致**：支出侧同一租户内**同号发票只算一次**（重复行不计入合计），
   与 `projects.summary()` 的规则相同；但这里不调用它，避免顺手在数据目录里建
   `projects.json`。
3. **不含邀请码**。码只在创建时显示一次，台账里只记账号名与状态，不记码。

## 用法

    # 服务器上（真实数据）
    CASH_DATA_DIR=/var/lib/cashlens python3 scripts/make_tenant_ledger.py --out ~/账号台账.xlsx

    # 本机（用拷下来的副本）
    python3 scripts/make_tenant_ledger.py --data-dir /tmp/cl_tenants --out 账号台账.xlsx
"""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timedelta
from pathlib import Path

import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app.services.state_engine.ledger import EventLedger  # noqa: E402

SHEET = "账号台账"
NOTE = "说明"
_MONEY = "#,##0.00"
_HEAD_FILL = PatternFill("solid", fgColor="EEF1F6")
_MANUAL_FILL = PatternFill("solid", fgColor="FFF8E6")   # 手填列：浅黄，一眼能分开

HEADERS = [
    ("账号", 16, "auto"),
    ("状态", 8, "auto"),
    ("创建日", 12, "auto"),
    ("最近使用", 12, "auto"),
    ("账本笔数", 10, "auto"),
    ("收入(元)", 13, "auto"),
    ("支出(元)", 13, "auto"),
    ("净额(元)", 13, "auto"),
    ("近30天笔数", 11, "auto"),
    ("最近一笔", 12, "auto"),
    ("反馈 / 用得怎么样", 34, "manual"),
    ("是否付费", 11, "manual"),
    ("备注", 26, "manual"),
]


def _load_tenants(data_dir: Path) -> list[dict]:
    p = data_dir / "tenants.json"
    if not p.exists():
        raise SystemExit(f"✗ 找不到 {p}（--data-dir 指到含 tenants.json 的目录）")
    try:
        rows = json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise SystemExit(f"✗ tenants.json 解析失败：{e}") from e
    return [r for r in rows if isinstance(r, dict) and r.get("slug")]


def _read_events(path: Path) -> list[dict]:
    """读一个租户的账本。

    ⚠️ 必须走 `EventLedger.load()`，不能自己逐行 json.loads —— 那个函数会
    ① 滤掉被 `type=void` 作废的事件、② 按 dedupe_key 只留先写者。
    自己读会把**已作废的账算进合计**（2026-10-20 实测：一次 void + 重记会让支出虚高）。
    """
    if not path.exists():
        return []
    return EventLedger(path).load()


def _invoice_no(ev: dict) -> str:
    extra = ev.get("extra") if isinstance(ev.get("extra"), dict) else {}
    return str(ev.get("invoice_no") or extra.get("invoice_no") or "").strip()


def _stats(events: list[dict], as_of: datetime) -> dict:
    """一个租户的概况。口径与工作台一致：同号发票在一个租户内只算一次支出。"""
    income = expense = 0
    count = 0
    seen: set[str] = set()
    latest = ""
    recent_30 = 0
    cutoff = as_of - timedelta(days=30)
    for ev in events:
        typ = str(ev.get("type") or "")
        ts = str(ev.get("ts") or "")
        amount = int(ev.get("amount_cents") or 0)
        if ts > latest:
            latest = ts
        try:
            when = datetime.fromisoformat(ts.replace("Z", "")) if ts else None
        except ValueError:
            when = None
        if when and when.replace(tzinfo=None) >= cutoff.replace(tzinfo=None):
            recent_30 += 1
        if typ in ("income", "refund"):
            income += amount
            count += 1
        elif typ == "expense":
            no = _invoice_no(ev)
            if no and no in seen:
                continue            # 同号重复票不计入合计（明细里仍留在账本）
            if no:
                seen.add(no)
            expense += amount
            count += 1
    return {
        "income": income, "expense": expense, "net": income - expense,
        "count": count, "recent_30": recent_30, "latest": latest[:10],
    }


def _load_manual(out_path: Path) -> dict[str, tuple[str, str, str]]:
    """读上一次生成的台账，把**手填列**按账号名带过来。

    不这么做的话，「重新生成」会把几周的反馈一次抹掉 —— 而台账的价值恰好在手填那几列。
    """
    if not out_path.exists():
        return {}
    try:
        wb = load_workbook(out_path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    except Exception:  # noqa: BLE001  文件坏了就当没有，别因此生不出新台账
        # ⚠️ 这条宽 except 曾把「未 import load_workbook」的 NameError 也吞掉，
        # 表现为「手填列莫名其妙没保留」——所以读取逻辑必须有测试兜着（见 tests）。
        return {}
    keep: dict[str, tuple[str, str, str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[0] or "").strip()
        if not name or name == "合计":
            continue
        keep[name] = (str(row[10] or ""), str(row[11] or ""), str(row[12] or ""))
    return keep


def build(data_dir: Path, out_path: Path) -> Path:
    as_of = datetime.now()
    tenants = _load_tenants(data_dir)
    manual = _load_manual(out_path)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    ws.append([h for h, _, _ in HEADERS])
    for i, (_, width, kind) in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = _MANUAL_FILL if kind == "manual" else _HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    totals = {"count": 0, "income": 0, "expense": 0, "net": 0}
    for r in sorted(tenants, key=lambda x: str(x.get("created_at") or "")):
        st = _stats(_read_events(data_dir / "tenants" / str(r["slug"]) / "finance_events.jsonl"), as_of)
        ws.append([
            str(r.get("label") or "（未命名）"),
            "在用" if r.get("active") else "停用",
            str(r.get("created_at") or "")[:10],
            str(r.get("last_seen") or "—")[:10],
            st["count"],
            st["income"] / 100,
            st["expense"] / 100,
            st["net"] / 100,
            st["recent_30"],
            st["latest"] or "—",
            *(manual.get(str(r.get("label") or "（未命名）"), ("", "", ""))),
        ])
        for k in totals:
            totals[k] += st[k]

    last = ws.max_row + 1
    ws.cell(row=last, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=last, column=5, value=totals["count"]).font = Font(bold=True)
    for col, key in ((6, "income"), (7, "expense"), (8, "net")):
        c = ws.cell(row=last, column=col, value=totals[key] / 100)
        c.font = Font(bold=True)
    for row in range(2, last + 1):
        for col in (6, 7, 8):
            ws.cell(row=row, column=col).number_format = _MONEY

    # ── 说明页：口径与「哪些列是自动的」写清楚，免得下次自己都分不清 ──
    ws2 = wb.create_sheet(NOTE)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 96
    lines = [
        ("生成时间", f"{as_of:%Y-%m-%d %H:%M}"),
        ("数据来源", str(data_dir)),
        ("账号数", f"{len(tenants)} 个（在用 {sum(1 for r in tenants if r.get('active'))}）"),
        ("", ""),
        ("自动列（灰）", "账号 / 状态 / 创建日 / 最近使用 / 账本笔数 / 收入 / 支出 / 净额 / 近30天笔数 / 最近一笔"),
        ("手填列（黄）", "反馈 / 是否付费 / 备注 —— 重新生成时按账号名保留旧值（脚本只写自动列）"),
        ("", ""),
        ("金额口径", "与工作台一致：收入含退款；支出侧同一账号内**同号发票只算一次**（重复票不进合计，账本里仍保留）"),
        ("笔数口径", "只数真正计入合计的笔（同号重复票不重复计）；「近30天笔数」按发生时间算"),
        ("不含什么", "不含邀请码明文（码只在创建时显示一次，库里只存哈希）"),
        ("重生成", "python3 scripts/make_tenant_ledger.py --data-dir <含 tenants.json 的目录> --out <文件名>.xlsx"),
    ]
    for k, v in lines:
        ws2.append([k, v])
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main() -> int:
    ap = argparse.ArgumentParser(description="生成邀请码账号台账（xlsx）")
    ap.add_argument("--data-dir", default="", help="含 tenants.json 的目录（默认取 CASH_DATA_DIR 或仓库 data/）")
    ap.add_argument("--out", default="", help="输出文件名（默认 账号台账-YYYYMMDD.xlsx）")
    args = ap.parse_args()

    if args.data_dir:
        data_dir = Path(args.data_dir).expanduser().resolve()
    else:
        import os
        env = os.environ.get("CASH_DATA_DIR")
        data_dir = Path(env).expanduser().resolve() if env else \
            (Path(__file__).resolve().parents[1] / "data")
    out = Path(args.out).expanduser() if args.out else \
        Path.cwd() / f"账号台账-{datetime.now():%Y%m%d}.xlsx"

    path = build(data_dir, out)
    print(f"✓ 已生成：{path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
